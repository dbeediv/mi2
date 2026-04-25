"""
ingestion.py  (v2 — Hybrid RAG)
================================
Upgrades over v1:
  - BM25 index built alongside FAISS (persisted per session)
  - Table-aware PDF parsing via pdfplumber — tables become readable text rows
  - Abstract/section priority tagging — abstract chunks always retrieved first
  - Reference/bibliography pages hard-rejected at parse time (never indexed)

Author: Senior AI Architect
"""

from __future__ import annotations

import hashlib
import json
import logging
import pickle
import re
import sqlite3
import uuid
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import List, Optional, Dict

import faiss
import fitz                            # PyMuPDF
import numpy as np
import pdfplumber
from docx import Document as DocxDocument
from rank_bm25 import BM25Okapi
from sentence_transformers import SentenceTransformer

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
logging.basicConfig(
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
EMBED_MODEL_NAME: str       = "BAAI/bge-base-en-v1.5"
CHUNK_SIZE: int             = 512
CHUNK_OVERLAP: int          = 64
SIMILARITY_THRESHOLD: float = 0.01  # RRF scores are small by nature (1/60 scale)
STORAGE_ROOT: Path          = Path("storage")
FAISS_DIR: Path             = STORAGE_ROOT / "faiss_store"
BM25_DIR: Path              = STORAGE_ROOT / "bm25_store"
METADATA_DB: Path           = STORAGE_ROOT / "metadata.db"
RAW_DOCS_DIR: Path          = STORAGE_ROOT / "raw_docs"

PRIORITY_SECTIONS = {"abstract", "introduction", "conclusion", "summary", "results"}
REJECT_SECTIONS   = {"references", "bibliography", "acknowledgement", "acknowledgments"}

for _dir in (FAISS_DIR, BM25_DIR, RAW_DOCS_DIR):
    _dir.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------------
# Data models
# ---------------------------------------------------------------------------
@dataclass
class ParsedChunk:
    chunk_id: str
    session_id: str
    text: str
    source_file: str
    page_number: int
    section: str
    char_start: int
    char_end: int
    is_table: bool = False
    is_priority: bool = False
    embedding: Optional[np.ndarray] = field(default=None, repr=False)

    def to_dict_no_embedding(self) -> dict:
        d = asdict(self)
        d.pop("embedding", None)
        return d


# ---------------------------------------------------------------------------
# SQLite metadata store
# ---------------------------------------------------------------------------
class MetadataStore:

    def __init__(self, db_path: Path = METADATA_DB) -> None:
        self.db_path = db_path
        self._init_db()

    def _init_db(self) -> None:
        with self._conn() as conn:
            conn.execute("""
                CREATE TABLE IF NOT EXISTS chunks (
                    chunk_id    TEXT PRIMARY KEY,
                    session_id  TEXT NOT NULL,
                    source_file TEXT NOT NULL,
                    page_number INTEGER NOT NULL,
                    section     TEXT,
                    char_start  INTEGER,
                    char_end    INTEGER,
                    text        TEXT NOT NULL,
                    is_table    INTEGER DEFAULT 0,
                    is_priority INTEGER DEFAULT 0,
                    created_at  DATETIME DEFAULT CURRENT_TIMESTAMP
                )
            """)
            conn.execute("CREATE INDEX IF NOT EXISTS idx_session ON chunks(session_id)")

    def _conn(self) -> sqlite3.Connection:
        return sqlite3.connect(str(self.db_path))

    def insert_chunks(self, chunks: List[ParsedChunk]) -> None:
        rows = [
            (
                c.chunk_id, c.session_id, c.source_file, c.page_number,
                c.section, c.char_start, c.char_end, c.text,
                int(c.is_table), int(c.is_priority)
            )
            for c in chunks
        ]
        with self._conn() as conn:
            conn.executemany(
                """INSERT OR REPLACE INTO chunks
                   (chunk_id, session_id, source_file, page_number,
                    section, char_start, char_end, text, is_table, is_priority)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                rows,
            )
        logger.info("Inserted %d chunks into metadata store.", len(rows))

    def fetch_chunks_by_ids(self, chunk_ids: List[str]) -> List[dict]:
        if not chunk_ids:
            return []
        placeholders = ",".join("?" * len(chunk_ids))
        with self._conn() as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                f"SELECT * FROM chunks WHERE chunk_id IN ({placeholders})",
                chunk_ids,
            ).fetchall()
        return [dict(r) for r in rows]

    def fetch_chunks_by_session(self, session_id: str, limit: int = 1200) -> List[dict]:
        with self._conn() as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                "SELECT * FROM chunks WHERE session_id = ? ORDER BY page_number ASC LIMIT ?",
                (session_id, limit),
            ).fetchall()
        return [dict(r) for r in rows]

    def fetch_priority_chunks(self, session_id: str) -> List[dict]:
        """Return abstract / conclusion / summary chunks."""
        with self._conn() as conn:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                "SELECT * FROM chunks WHERE session_id = ? AND is_priority = 1 ORDER BY page_number ASC",
                (session_id,),
            ).fetchall()
        return [dict(r) for r in rows]

    def delete_session(self, session_id: str) -> None:
        with self._conn() as conn:
            conn.execute("DELETE FROM chunks WHERE session_id = ?", (session_id,))
        logger.info("Deleted session %s from metadata store.", session_id)


# ---------------------------------------------------------------------------
# BM25 Index Manager
# ---------------------------------------------------------------------------
class BM25IndexManager:
    """
    Sparse keyword index (BM25Okapi) per session.
    Complements FAISS: finds exact matches for numbers, acronyms, model names
    that dense embeddings often miss.
    """

    def __init__(self, index_dir: Path = BM25_DIR) -> None:
        self.index_dir = index_dir
        self._indexes: Dict[str, BM25Okapi]       = {}
        self._id_maps: Dict[str, List[str]]        = {}
        self._corpus:  Dict[str, List[List[str]]]  = {}

    @staticmethod
    def _tokenize(text: str) -> List[str]:
        return re.findall(r"[a-zA-Z0-9]+", text.lower())

    def build_or_update(self, session_id: str, chunks: List[ParsedChunk]) -> None:
        existing_ids    = self._id_maps.get(session_id, [])
        existing_corpus = self._corpus.get(session_id, [])

        new_ids    = [c.chunk_id for c in chunks]
        new_corpus = [self._tokenize(c.text) for c in chunks]

        all_ids    = existing_ids + new_ids
        all_corpus = existing_corpus + new_corpus

        index = BM25Okapi(all_corpus)
        self._indexes[session_id] = index
        self._id_maps[session_id] = all_ids
        self._corpus[session_id]  = all_corpus

        self._persist(session_id)
        logger.info("BM25 updated. Session: %s | Total docs: %d", session_id, len(all_ids))

    def search(self, session_id: str, query: str, top_k: int = 6) -> List[tuple[str, float]]:
        index, id_map = self._load_if_needed(session_id)
        if index is None or not id_map:
            return []

        tokens = self._tokenize(query)
        if not tokens:
            return []

        scores = index.get_scores(tokens)
        max_score = scores.max()
        if max_score == 0:
            return []

        norm_scores   = scores / max_score
        top_indices   = np.argsort(norm_scores)[::-1][:top_k]
        return [(id_map[i], float(norm_scores[i])) for i in top_indices if norm_scores[i] > 0]

    def session_exists(self, session_id: str) -> bool:
        return self._pickle_path(session_id).exists()

    def delete_session(self, session_id: str) -> None:
        p = self._pickle_path(session_id)
        if p.exists():
            p.unlink()
        self._indexes.pop(session_id, None)
        self._id_maps.pop(session_id, None)
        self._corpus.pop(session_id, None)

    def _load_if_needed(self, session_id: str) -> tuple[Optional[BM25Okapi], List[str]]:
        if session_id in self._indexes:
            return self._indexes[session_id], self._id_maps[session_id]
        p = self._pickle_path(session_id)
        if not p.exists():
            return None, []
        with open(p, "rb") as f:
            data = pickle.load(f)
        self._indexes[session_id] = data["index"]
        self._id_maps[session_id] = data["id_map"]
        self._corpus[session_id]  = data["corpus"]
        logger.info("Loaded BM25 from disk for session %s.", session_id)
        return self._indexes[session_id], self._id_maps[session_id]

    def _persist(self, session_id: str) -> None:
        with open(self._pickle_path(session_id), "wb") as f:
            pickle.dump(
                {"index": self._indexes[session_id],
                 "id_map": self._id_maps[session_id],
                 "corpus": self._corpus[session_id]},
                f,
            )

    def _pickle_path(self, session_id: str) -> Path:
        return self.index_dir / f"{session_id}.bm25.pkl"


# ---------------------------------------------------------------------------
# Document parser
# ---------------------------------------------------------------------------
class DocumentParser:
    """Two-pass PDF (pdfplumber tables + PyMuPDF text), DOCX, TXT."""

    def parse(self, file_path: Path) -> List[tuple[str, int, str, bool]]:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return self._parse_pdf(file_path)
        elif suffix in (".docx", ".doc"):
            return self._parse_docx(file_path)
        elif suffix == ".txt":
            return self._parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {suffix}")

    def _parse_pdf(self, path: Path) -> List[tuple[str, int, str, bool]]:
        results: List[tuple[str, int, str, bool]] = []

        # Pass 1 — tables (pdfplumber)
        table_pages: Dict[int, List[str]] = {}
        try:
            with pdfplumber.open(str(path)) as pdf:
                for page in pdf.pages:
                    pnum = page.page_number
                    for table in (page.extract_tables() or []):
                        if not table:
                            continue
                        headers = [str(c).strip() if c else "" for c in table[0]]
                        rows_text = []
                        for row in table[1:]:
                            cells = [str(c).strip() if c else "" for c in row]
                            row_str = " | ".join(
                                f"{h}: {v}" for h, v in zip(headers, cells) if v
                            )
                            if row_str:
                                rows_text.append(row_str)
                        if rows_text:
                            table_pages.setdefault(pnum, []).extend(rows_text)
        except Exception as e:
            logger.warning("pdfplumber failed: %s", e)

        for pnum, rows in table_pages.items():
            results.append(("\n".join(rows), pnum, "table", True))

        # Pass 2 — body text (PyMuPDF)
        doc = fitz.open(str(path))
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if not text:
                continue

            lines = [l.strip() for l in text.split("\n") if l.strip()]
            first_line = lines[0] if lines else ""
            section = first_line if len(first_line) < 100 else ""
            section_lower = section.lower()

            # Hard-reject reference/bibliography pages
            if any(k in section_lower for k in REJECT_SECTIONS):
                continue
            # Hard-reject pages that are mostly [N] citation entries
            citation_count = len(re.findall(r"\[\d+\]", text))
            if citation_count > 8 and len(text) < 3500:
                continue

            results.append((text, page_num, section, False))

        doc.close()
        logger.info("PDF parsed: %d blocks (incl. %d table blocks) from %s",
                    len(results), len(table_pages), path.name)
        return results

    def _parse_docx(self, path: Path) -> List[tuple[str, int, str, bool]]:
        doc = DocxDocument(str(path))
        blocks: List[tuple[str, int, str, bool]] = []
        current_section = ""
        buffer: List[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name or ""
            if style_name.startswith("Heading"):
                if buffer:
                    if not any(k in current_section.lower() for k in REJECT_SECTIONS):
                        blocks.append((" ".join(buffer), 0, current_section, False))
                    buffer = []
                current_section = text
            else:
                buffer.append(text)

        if buffer and not any(k in current_section.lower() for k in REJECT_SECTIONS):
            blocks.append((" ".join(buffer), 0, current_section, False))

        return blocks

    def _parse_txt(self, path: Path) -> List[tuple[str, int, str, bool]]:
        text = path.read_text(encoding="utf-8", errors="replace").strip()
        return [(text, 0, "", False)]


# ---------------------------------------------------------------------------
# Chunker
# ---------------------------------------------------------------------------
class TextChunker:

    def __init__(self, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> None:
        self.chunk_chars   = chunk_size * 4
        self.overlap_chars = overlap * 4

    def chunk(
        self,
        text: str,
        session_id: str,
        source_file: str,
        page_number: int,
        section: str,
        is_table: bool = False,
    ) -> List[ParsedChunk]:
        if is_table:
            cid = self._make_id(session_id, source_file, 0, section)
            return [ParsedChunk(
                chunk_id=cid, session_id=session_id, text=text.strip(),
                source_file=source_file, page_number=page_number, section=section,
                char_start=0, char_end=len(text), is_table=True, is_priority=False,
            )]

        section_lower = section.lower()
        is_priority   = any(k in section_lower for k in PRIORITY_SECTIONS)

        chunks: List[ParsedChunk] = []
        start = 0
        text_len = len(text)

        while start < text_len:
            end = min(start + self.chunk_chars, text_len)
            if end < text_len:
                boundary = text.rfind(". ", start, end)
                if boundary != -1 and boundary > start + self.overlap_chars:
                    end = boundary + 1
            chunk_text = text[start:end].strip()
            if chunk_text:
                cid = self._make_id(session_id, source_file, start, section)
                chunks.append(ParsedChunk(
                    chunk_id=cid, session_id=session_id, text=chunk_text,
                    source_file=source_file, page_number=page_number, section=section,
                    char_start=start, char_end=end, is_table=False, is_priority=is_priority,
                ))
            start = end - self.overlap_chars
            if start >= text_len or end == text_len:
                break

        return chunks

    @staticmethod
    def _make_id(session_id: str, source_file: str, offset: int, section: str) -> str:
        raw = f"{session_id}::{source_file}::{section}::{offset}"
        return hashlib.sha256(raw.encode()).hexdigest()[:24]


# ---------------------------------------------------------------------------
# Embedding model
# ---------------------------------------------------------------------------
class EmbeddingModel:
    QUERY_PREFIX = "Represent this sentence for searching relevant passages: "

    def __init__(self, model_name: str = EMBED_MODEL_NAME) -> None:
        logger.info("Loading embedding model: %s", model_name)
        self.model = SentenceTransformer(model_name)
        self.dim: int = self.model.get_sentence_embedding_dimension()
        logger.info("Embedding model loaded. Dimension: %d", self.dim)

    def embed_documents(self, texts: List[str]) -> np.ndarray:
        return np.array(
            self.model.encode(texts, batch_size=32, normalize_embeddings=True,
                              show_progress_bar=len(texts) > 50),
            dtype=np.float32,
        )

    def embed_query(self, query: str) -> np.ndarray:
        return np.array(
            self.model.encode(self.QUERY_PREFIX + query, normalize_embeddings=True),
            dtype=np.float32,
        )


# ---------------------------------------------------------------------------
# FAISS Index Manager
# ---------------------------------------------------------------------------
class FAISSIndexManager:

    def __init__(self, index_dir: Path = FAISS_DIR) -> None:
        self.index_dir = index_dir
        self._indexes: dict[str, faiss.IndexFlatIP] = {}
        self._id_maps: dict[str, List[str]]          = {}

    def build_or_update(self, session_id: str, chunks: List[ParsedChunk], dim: int) -> None:
        index  = self._get_or_create(session_id, dim)
        id_map = self._id_maps.setdefault(session_id, [])
        vectors = np.stack([c.embedding for c in chunks]).astype(np.float32)
        index.add(vectors)
        id_map.extend(c.chunk_id for c in chunks)
        self._persist(session_id)
        logger.info("FAISS updated. Session: %s | Total: %d", session_id, index.ntotal)

    def search(self, session_id: str, query_vector: np.ndarray, top_k: int = 6) -> List[tuple[str, float]]:
        index, id_map = self._load_if_needed(session_id)
        if index is None or index.ntotal == 0:
            return []
        k = min(top_k, index.ntotal)
        scores, indices = index.search(query_vector.reshape(1, -1), k)
        return [(id_map[i], float(s)) for s, i in zip(scores[0], indices[0]) if i != -1]

    def session_exists(self, session_id: str) -> bool:
        return self._index_path(session_id).exists()

    def delete_session(self, session_id: str) -> None:
        for ext in (".index", ".ids.json"):
            p = self.index_dir / f"{session_id}{ext}"
            if p.exists():
                p.unlink()
        self._indexes.pop(session_id, None)
        self._id_maps.pop(session_id, None)

    def _get_or_create(self, session_id: str, dim: int) -> faiss.IndexFlatIP:
        if session_id not in self._indexes:
            self._indexes[session_id] = faiss.IndexFlatIP(dim)
        return self._indexes[session_id]

    def _load_if_needed(self, session_id: str) -> tuple[Optional[faiss.IndexFlatIP], List[str]]:
        if session_id in self._indexes:
            return self._indexes[session_id], self._id_maps[session_id]
        ip = self._index_path(session_id)
        if not ip.exists():
            return None, []
        index = faiss.read_index(str(ip))
        with open(ip.with_suffix(".ids.json")) as f:
            id_map = json.load(f)
        self._indexes[session_id] = index
        self._id_maps[session_id] = id_map
        return index, id_map

    def _persist(self, session_id: str) -> None:
        ip = self._index_path(session_id)
        faiss.write_index(self._indexes[session_id], str(ip))
        with open(ip.with_suffix(".ids.json"), "w") as f:
            json.dump(self._id_maps[session_id], f)

    def _index_path(self, session_id: str) -> Path:
        return self.index_dir / f"{session_id}.index"


# ---------------------------------------------------------------------------
# Ingestion Pipeline
# ---------------------------------------------------------------------------
class IngestionPipeline:

    def __init__(self) -> None:
        self.parser         = DocumentParser()
        self.chunker        = TextChunker()
        self.embedder       = EmbeddingModel()
        self.faiss_mgr      = FAISSIndexManager()
        self.bm25_mgr       = BM25IndexManager()
        self.metadata_store = MetadataStore()

    def ingest(self, file_path: Path, session_id: str) -> int:
        logger.info("Ingestion | session=%s | file=%s", session_id, file_path.name)

        try:
            raw_blocks = self.parser.parse(file_path)
        except Exception as exc:
            raise RuntimeError(f"Parse failed: {exc}") from exc

        if not raw_blocks:
            return 0

        all_chunks: List[ParsedChunk] = []
        for text, page_num, section, is_table in raw_blocks:
            all_chunks.extend(self.chunker.chunk(
                text=text, session_id=session_id, source_file=file_path.name,
                page_number=page_num, section=section, is_table=is_table,
            ))

        if not all_chunks:
            return 0

        # Embed
        vectors = self.embedder.embed_documents([c.text for c in all_chunks])
        for chunk, vec in zip(all_chunks, vectors):
            chunk.embedding = vec

        # Index: FAISS (dense) + BM25 (sparse)
        self.faiss_mgr.build_or_update(session_id, all_chunks, self.embedder.dim)
        self.bm25_mgr.build_or_update(session_id, all_chunks)
        self.metadata_store.insert_chunks(all_chunks)

        logger.info(
            "Done | chunks=%d | priority=%d | tables=%d",
            len(all_chunks),
            sum(c.is_priority for c in all_chunks),
            sum(c.is_table for c in all_chunks),
        )
        return len(all_chunks)

    def ingest_many(self, file_paths: List[Path], session_id: str) -> dict[str, int]:
        results: dict[str, int] = {}
        for fp in file_paths:
            try:
                results[fp.name] = self.ingest(fp, session_id)
            except (ValueError, RuntimeError) as exc:
                logger.error("Failed to ingest %s: %s", fp.name, exc)
                results[fp.name] = -1
        return results

    def generate_session_id(self) -> str:
        return str(uuid.uuid4())